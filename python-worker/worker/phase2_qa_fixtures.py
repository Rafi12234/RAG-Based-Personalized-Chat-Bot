"""Phase 2.12: generate *synthetic* PDF, scanned PDF, DOCX and Bengali TXT QA files.

Files are created on the local filesystem only. Existing fixture files are
NEVER overwritten. No MySQL, Redis, Qdrant, RustFS or Celery operations.

Run from python-worker:
  python -m worker.phase2_qa_fixtures --output-dir ..\\qa-fixtures
  python -m worker.phase2_qa_fixtures --output-dir ..\\qa-1000p --native-pdf-pages 1000
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

MAX_NATIVE_PAGES = 1000
MAX_SCANNED_PAGES = 10
MAX_DOCX_PARAGRAPHS = 5000
MAX_TXT_LINES = 100_000

_FIXTURES = (
    "qa_native.pdf", "qa_scanned.pdf", "qa_structured.docx",
    "qa_bengali.txt", "qa_manifest.json",
)


def _int_range(value: int, name: str, low: int, high: int) -> None:
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f"{name} must be {low}..{high}")


def _create_native_pdf(path: Path, pages: int) -> None:
    import pymupdf
    pdf = pymupdf.open()
    try:
        for number in range(1, pages + 1):
            page = pdf.new_page(width=612, height=792)
            title = f"DOCMIND QA NATIVE PAGE {number:04d} OF {pages:04d}"
            page.insert_text((48, 70), title, fontsize=17)
            body = (
                "Document ingestion quality assurance example. "
                "Each page carries unique source metadata and sufficient searchable text. "
                f"Source page number {number}. The content is synthetic, non-sensitive. "
            ) * 3
            overflow = page.insert_textbox(pymupdf.Rect(48, 115, 565, 700), body, fontsize=12)
            if overflow < 0:
                raise RuntimeError("Synthetic native PDF text did not fit")
        pdf.save(str(path), garbage=3, deflate=True)
    finally:
        pdf.close()


def _create_scanned_pdf(path: Path, pages: int) -> None:
    """Make image-only pages, not searchable PDF text (real OCR smoke fixture)."""
    import pymupdf
    scanned = pymupdf.open()
    try:
        for number in range(1, pages + 1):
            temporary = pymupdf.open()
            try:
                source = temporary.new_page(width=612, height=792)
                source.insert_text((42, 110), "DOCMIND OCR VALIDATION", fontsize=25)
                source.insert_text((42, 165), f"SCANNED TEST PAGE {number}", fontsize=21)
                source.insert_text((42, 220), "Recognize this English text with Tesseract OCR.", fontsize=15)
                png_bytes = source.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).tobytes("png")
            finally:
                temporary.close()
            target = scanned.new_page(width=612, height=792)
            target.insert_image(target.rect, stream=png_bytes)
        scanned.save(str(path), garbage=3, deflate=True)
    finally:
        scanned.close()


def _create_docx(path: Path, paragraphs: int) -> None:
    from docx import Document
    document = Document()
    document.add_heading("DocMind AI - synthetic structured document", 0)
    for number in range(1, paragraphs + 1):
        document.add_paragraph(
            f"QA paragraph {number}. This sample validates document order, "
            "structured extraction, Unicode text and safe chunking. "
            "Everything here is fabricated for testing."
        )
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Test Case"
    table.rows[0].cells[1].text = "Expected Output"
    for key, value in (("DOCX extraction", "Paragraphs in order"),
                       ("DOCX tables", "Table-row segments")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    document.save(path)


def _create_bengali_txt(path: Path, lines: int) -> None:
    phrase = "ডকমাইন্ড এআই একটি নথিভিত্তিক প্রশ্নোত্তর ব্যবস্থা। এটি বাংলা লেখা পরীক্ষা করে।"
    with path.open("x", encoding="utf-8", newline="\n") as output:
        for number in range(1, lines + 1):
            output.write(f"পরীক্ষা {number}: {phrase}\n")


def generate_fixture_set(output_dir: str | Path, *, native_pdf_pages: int = 20,
                         scanned_pdf_pages: int = 2, docx_paragraphs: int = 50,
                         txt_lines: int = 1000) -> dict:
    """Produce a deterministic-sized synthetic corpus; refuse existing files."""
    _int_range(native_pdf_pages, "native_pdf_pages", 1, MAX_NATIVE_PAGES)
    _int_range(scanned_pdf_pages, "scanned_pdf_pages", 1, MAX_SCANNED_PAGES)
    _int_range(docx_paragraphs, "docx_paragraphs", 1, MAX_DOCX_PARAGRAPHS)
    _int_range(txt_lines, "txt_lines", 1, MAX_TXT_LINES)
    out = Path(output_dir).expanduser().resolve()
    if out.exists() and not out.is_dir():
        raise ValueError("Output location is not a directory")
    existing = [name for name in _FIXTURES if (out / name).exists()]
    if existing:
        raise FileExistsError("Existing QA fixture files would be overwritten")
    out.mkdir(parents=True, exist_ok=True)
    _create_native_pdf(out / _FIXTURES[0], native_pdf_pages)
    _create_scanned_pdf(out / _FIXTURES[1], scanned_pdf_pages)
    _create_docx(out / _FIXTURES[2], docx_paragraphs)
    _create_bengali_txt(out / _FIXTURES[3], txt_lines)
    report = {
        "fixture_directory": str(out),
        "native_pdf_pages": native_pdf_pages,
        "scanned_pdf_pages": scanned_pdf_pages,
        "docx_paragraphs": docx_paragraphs,
        "bengali_txt_lines": txt_lines,
        "files": [
            {"name": name, "bytes": (out / name).stat().st_size}
            for name in _FIXTURES[:-1]
        ],
        "synthetic_data_only": True,
        "mysql_modified": False,
        "qdrant_modified": False,
    }
    (out / _FIXTURES[-1]).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--native-pdf-pages", type=int, default=20)
    parser.add_argument("--scanned-pdf-pages", type=int, default=2)
    parser.add_argument("--docx-paragraphs", type=int, default=50)
    parser.add_argument("--txt-lines", type=int, default=1000)
    a = parser.parse_args(argv)
    try:
        report = generate_fixture_set(
            a.output_dir, native_pdf_pages=a.native_pdf_pages,
            scanned_pdf_pages=a.scanned_pdf_pages,
            docx_paragraphs=a.docx_paragraphs, txt_lines=a.txt_lines,
        )
        print(json.dumps(report, indent=2, ensure_ascii=False))
        return 0
    except Exception as exc:
        print("Fixture generation refused/failed: " + type(exc).__name__)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
