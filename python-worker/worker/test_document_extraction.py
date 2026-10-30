"""Offline file extraction tests: no MySQL, RustFS, Redis or Qdrant calls."""

import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import Mock

import pymupdf
from docx import Document

from .document_extraction import (
    DocumentExtractionError,
    iter_extracted_segments,
    summarize_extraction,
)
from .storage_extraction_probe import download_into


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="docmind-unit-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_pdf_page_numbers_and_text(self):
        path = self.root / "demo.pdf"
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 72), "First page contains readable document content")
            page = pdf.new_page()
            page.insert_text((72, 72), "Second page contains another paragraph of document text")
            pdf.save(str(path))
        segments = list(iter_extracted_segments(path, "demo.pdf"))
        self.assertEqual([s.page_number for s in segments], [1, 2])
        self.assertIn("First page", segments[0].text)
        summary = summarize_extraction(path, "demo.pdf")
        self.assertEqual(summary.pages, 2)
        self.assertEqual(summary.potential_ocr_pages, ())

    def test_pdf_empty_page_flagged_for_ocr(self):
        path = self.root / "scanned.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(str(path))
        summary = summarize_extraction(path, "scanned.pdf")
        self.assertEqual(summary.potential_ocr_pages, (1,))

    def test_docx_preserves_paragraph_and_table_order(self):
        path = self.root / "sample.docx"
        doc = Document()
        doc.add_paragraph("Introduction")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Column A"
        table.rows[0].cells[1].text = "Column B"
        doc.add_paragraph("Final paragraph")
        doc.save(path)
        segments = list(iter_extracted_segments(path, "sample.docx"))
        self.assertEqual([s.kind for s in segments], ["docx_paragraph", "docx_table_row", "docx_paragraph"])
        self.assertEqual(segments[1].text, "Column A | Column B")

    def test_txt_utf8_bom(self):
        path = self.root / "sample.txt"
        path.write_bytes(b"\xef\xbb\xbfHello, DocMind AI")
        segments = list(iter_extracted_segments(path, "sample.txt"))
        self.assertEqual("".join(s.text for s in segments), "Hello, DocMind AI")

    def test_txt_invalid_utf8_rejected(self):
        path = self.root / "bad.txt"
        path.write_bytes(b"\xff\xfe\x01")
        with self.assertRaises(DocumentExtractionError):
            list(iter_extracted_segments(path, "bad.txt"))

    def test_unsupported_extension_rejected(self):
        path = self.root / "file.exe"
        path.write_text("data")
        with self.assertRaises(DocumentExtractionError):
            list(iter_extracted_segments(path, "file.exe"))

    def test_streaming_download_validates_length(self):
        mock_client = Mock()
        body = BytesIO(b"hello world")
        mock_client.get_object.return_value = {"Body": body, "ContentLength": 11}
        dest = self.root / "output.txt"
        self.assertEqual(download_into(mock_client, "bucket", "key", dest, 11), 11)
        self.assertEqual(dest.read_bytes(), b"hello world")
        self.assertTrue(body.closed)

    def test_streaming_download_rejects_wrong_length(self):
        mock_client = Mock()
        body = BytesIO(b"hello")
        mock_client.get_object.return_value = {"Body": body, "ContentLength": 5}
        with self.assertRaises(DocumentExtractionError):
            download_into(mock_client, "bucket", "key", self.root / "file.txt", 6)
        self.assertTrue(body.closed)


if __name__ == "__main__":
    unittest.main()
